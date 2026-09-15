"""Portable ingestion contract adapted from DocCluster's parser/chunker layer."""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class ParsedDocument:
    filename: str
    content: str
    pages: tuple[str, ...]


def parse_document(filename: str, payload: bytes) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        content = payload.decode("utf-8", errors="replace")
        return ParsedDocument(filename, content, tuple(content.split("\f")))
    if suffix == ".pdf":
        from pypdf import PdfReader
        import io
        pages = tuple((page.extract_text() or "") for page in PdfReader(io.BytesIO(payload)).pages)
        return ParsedDocument(filename, "\n\n".join(pages), pages)
    if suffix == ".docx":
        from docx import Document
        import io
        document = Document(io.BytesIO(payload))
        content = "\n".join(p.text for p in document.paragraphs if p.text.strip())
        return ParsedDocument(filename, content, (content,))
    raise ValueError(f"unsupported file type: {suffix}")


def content_hash(content: str) -> str:
    return hashlib.sha256(content.encode("utf-8")).hexdigest()


def chunk_document(document: ParsedDocument, target_words: int = 180) -> list[dict]:
    chunks: list[dict] = []
    for page_number, page in enumerate(document.pages, start=1):
        words = re.findall(r"\S+", page)
        for offset in range(0, len(words), target_words):
            text = " ".join(words[offset:offset + target_words]).strip()
            if text:
                from app.core.providers import embed_text
                vector = embed_text(text)
                chunks.append({"text": text, "page": page_number, "ordinal": len(chunks), "heading": "", "embedding_json": json.dumps(vector), "embedding": vector})
    return chunks
